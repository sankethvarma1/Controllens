"""
Document Ingestion Pipeline for CONTROLLENS
Handles: parsing, cleaning, section detection, chunking, embedding, pgvector storage
"""
import os
import re
import hashlib
from datetime import datetime
from typing import List, Dict, Any, Optional, Tuple
from dataclasses import dataclass
from pathlib import Path
import json

# Document processing
try:
    import pdfplumber
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# Text processing
from langchain_text_splitters import RecursiveCharacterTextSplitter
# NOTE: SentenceTransformer is imported lazily inside EmbeddingGenerator.model
# so that backend startup and non-ML endpoints never require torch/transformers.

# Database
from sqlalchemy.orm import Session
from app.db.base import Document, DocumentChunk, SessionLocal


@dataclass
class DocumentSection:
    """Represents a detected section in a document."""
    title: str
    number: Optional[str]
    content: str
    page_number: int
    char_start: int
    char_end: int
    level: int  # Heading level (1=top, 2=subsection, etc.)


@dataclass
class ProcessedChunk:
    """Represents a processed text chunk ready for embedding."""
    content: str
    section_title: str
    section_number: Optional[str]
    page_number: int
    char_start: int
    char_end: int
    token_count: int
    metadata: Dict[str, Any]


class DocumentParser:
    """Parse documents into structured text with page information."""
    
    def __init__(self):
        self.supported_types = {
            'application/pdf': self._parse_pdf,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document': self._parse_docx,
            'text/plain': self._parse_text
        }
    
    def parse(self, file_path: str, mime_type: str) -> List[Tuple[str, int]]:
        """Parse document and return list of (text, page_number) tuples."""
        parser = self.supported_types.get(mime_type)
        if not parser:
            raise ValueError(f"Unsupported mime type: {mime_type}")
        return parser(file_path)
    
    def _parse_pdf(self, file_path: str) -> List[Tuple[str, int]]:
        """Extract text from PDF with page numbers."""
        if not PDF_AVAILABLE:
            raise RuntimeError("pdfplumber not installed")
        
        pages = []
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                pages.append((text, i + 1))
        return pages
    
    def _parse_docx(self, file_path: str) -> List[Tuple[str, int]]:
        """Extract text from DOCX (single page)."""
        if not DOCX_AVAILABLE:
            raise RuntimeError("python-docx not installed")
        
        doc = DocxDocument(file_path)
        text = "\n".join([p.text for p in doc.paragraphs])
        return [(text, 1)]
    
    def _parse_text(self, file_path: str) -> List[Tuple[str, int]]:
        """Read plain text file."""
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
        return [(text, 1)]


class TextCleaner:
    """Clean and normalize extracted text."""
    
    def __init__(self):
        # Patterns to clean
        self.patterns = [
            (r'\x00', ''),  # Null bytes
            (r'\ufeff', ''),  # BOM
            (r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', ''),  # Control chars
            (r' {3,}', '  '),  # Multiple spaces
            (r'\n{3,}', '\n\n'),  # Multiple newlines
            (r'\t+', ' '),  # Tabs to spaces
            (r'\r\n?', '\n'),  # Normalize line endings
        ]
    
    def clean(self, text: str) -> str:
        """Apply cleaning patterns."""
        for pattern, replacement in self.patterns:
            text = re.sub(pattern, replacement, text)
        return text.strip()
    
    def normalize_whitespace(self, text: str) -> str:
        """Normalize whitespace while preserving paragraph structure."""
        # Split into paragraphs
        paragraphs = text.split('\n\n')
        # Clean each paragraph
        cleaned = [re.sub(r'\s+', ' ', p).strip() for p in paragraphs]
        # Rejoin with double newline
        return '\n\n'.join(p for p in cleaned if p)


class SectionDetector:
    """Detect document sections using heading patterns."""
    
    def __init__(self):
        # Common heading patterns
        self.heading_patterns = [
            # Numbered sections: 1., 1.1, 1.1.1, Article 1, Section 1
            (r'^(Article|Section|Chapter|Part|Annex|Appendix)\s+(\d+[\.\d]*)\s*[:\-]?\s*(.+)$', 1),
            (r'^(\d+[\.\d]*)\s+(.+)$', 2),  # 1.1 Title
            # All caps headings
            (r'^([A-Z][A-Z\s]{3,})$', 1),
            # Title case headings (short lines)
            (r'^([A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,5})$', 2),
            # Markdown style
            (r'^(#{1,3})\s+(.+)$', None),  # Level from # count
        ]
    
    def detect_sections(self, pages: List[Tuple[str, int]]) -> List[DocumentSection]:
        """Detect sections across all pages."""
        sections = []
        current_section = None
        char_offset = 0
        
        for page_text, page_num in pages:
            lines = page_text.split('\n')
            page_content = ""
            
            for line in lines:
                line = line.strip()
                if not line:
                    page_content += '\n'
                    char_offset += 1
                    continue
                
                # Check for heading
                heading_match = self._match_heading(line)
                if heading_match:
                    # Save previous section
                    if current_section:
                        current_section.char_end = char_offset
                        current_section.content = page_content.strip()
                        sections.append(current_section)
                    
                    # Start new section
                    level, number, title = heading_match
                    current_section = DocumentSection(
                        title=title,
                        number=number,
                        content="",
                        page_number=page_num,
                        char_start=char_offset,
                        char_end=char_offset,
                        level=level
                    )
                    page_content += line + '\n'
                    char_offset += len(line) + 1
                else:
                    page_content += line + '\n'
                    char_offset += len(line) + 1
            
            # Handle page break
            if current_section and page_content.strip():
                current_section.content += page_content
        
        # Add last section
        if current_section:
            current_section.char_end = char_offset
            sections.append(current_section)
        
        # If no sections detected, treat whole document as one section
        if not sections and pages:
            full_text = "\n".join(p[0] for p in pages)
            sections.append(DocumentSection(
                title="Document",
                number=None,
                content=full_text,
                page_number=1,
                char_start=0,
                char_end=len(full_text),
                level=1
            ))
        
        return sections
    
    def _match_heading(self, line: str) -> Optional[Tuple[int, Optional[str], str]]:
        """Match line against heading patterns. Returns (level, number, title)."""
        for pattern, default_level in self.heading_patterns:
            match = re.match(pattern, line, re.IGNORECASE)
            if match:
                groups = match.groups()
                if default_level is None:  # Markdown
                    level = len(groups[0])
                    title = groups[1]
                    number = None
                elif len(groups) == 3:  # Article 1: Title
                    level = default_level
                    number = groups[1]
                    title = groups[2]
                elif len(groups) == 2:  # 1.1 Title
                    level = default_level
                    number = groups[0]
                    title = groups[1]
                else:  # Single group
                    level = default_level
                    number = None
                    title = groups[0]
                
                return (level, number, title.strip())
        return None


class Chunker:
    """Split text into overlapping chunks for embedding."""
    
    def __init__(self, chunk_size: int = 500, chunk_overlap: int = 100):
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            length_function=len,
            separators=["\n\n", "\n", ". ", " ", ""]
        )
    
    def chunk_section(self, section: DocumentSection) -> List[ProcessedChunk]:
        """Split a section into chunks."""
        chunks = self.splitter.split_text(section.content)
        
        processed = []
        char_pos = section.char_start
        
        for i, chunk_text in enumerate(chunks):
            token_count = len(chunk_text) // 4  # Rough estimate
            
            processed.append(ProcessedChunk(
                content=chunk_text,
                section_title=section.title,
                section_number=section.number,
                page_number=section.page_number,
                char_start=char_pos,
                char_end=char_pos + len(chunk_text),
                token_count=token_count,
                metadata={
                    "section_level": section.level,
                    "chunk_index": i
                }
            ))
            char_pos += len(chunk_text) - self.splitter._chunk_overlap
        
        return processed


class EmbeddingGenerator:
    """Generate embeddings using sentence-transformers."""
    
    def __init__(self, model_name: str = "all-MiniLM-L6-v2"):
        self.model_name = model_name
        self._model = None
    
    @property
    def model(self):
        if self._model is None:
            from sentence_transformers import SentenceTransformer
            self._model = SentenceTransformer(self.model_name)
        return self._model
    
    def embed(self, texts: List[str]) -> List[List[float]]:
        """Generate embeddings for list of texts."""
        embeddings = self.model.encode(texts, convert_to_numpy=True, show_progress_bar=False)
        return embeddings.tolist()
    
    def embed_single(self, text: str) -> List[float]:
        """Generate embedding for single text."""
        return self.embed([text])[0]


class DocumentIngestionPipeline:
    """Complete document ingestion pipeline."""
    
    def __init__(self, db: Session):
        self.db = db
        self.parser = DocumentParser()
        self.cleaner = TextCleaner()
        self.section_detector = SectionDetector()
        self.chunker = Chunker()
        self.embedder = EmbeddingGenerator()
    
    def ingest(self, file_path: str, title: str, source: str, 
               document_type: str, publication_date: datetime = None,
               mime_type: str = None) -> Document:
        """Run full ingestion pipeline."""
        
        # Detect mime type if not provided
        if mime_type is None:
            mime_type = self._detect_mime_type(file_path)
        
        # Calculate file hash
        file_hash = self._calculate_file_hash(file_path)
        
        # Check for duplicate
        existing = self.db.query(Document).filter(Document.file_hash == file_hash).first()
        if existing:
            print(f"Document already exists: {existing.id}")
            return existing
        
        # Create document record
        document = Document(
            title=title,
            source=source,
            document_type=document_type,
            file_path=file_path,
            file_hash=file_hash,
            mime_type=mime_type,
            publication_date=publication_date.date() if publication_date else None,
            ingestion_status="processing"
        )
        self.db.add(document)
        self.db.flush()
        
        try:
            # Parse
            print(f"  Parsing {file_path}...")
            pages = self.parser.parse(file_path, mime_type)
            document.page_count = len(pages)
            
            # Clean
            print(f"  Cleaning text...")
            cleaned_pages = [(self.cleaner.clean(text), pn) for text, pn in pages]
            
            # Detect sections
            print(f"  Detecting sections...")
            sections = self.section_detector.detect_sections(cleaned_pages)
            
            # Chunk
            print(f"  Chunking...")
            all_chunks = []
            for section in sections:
                chunks = self.chunker.chunk_section(section)
                all_chunks.extend(chunks)
            
            # Generate embeddings
            print(f"  Generating embeddings for {len(all_chunks)} chunks...")
            chunk_texts = [c.content for c in all_chunks]
            embeddings = self.embedder.embed(chunk_texts)
            
            # Store chunks
            print(f"  Storing chunks...")
            for i, (chunk, embedding) in enumerate(zip(all_chunks, embeddings)):
                doc_chunk = DocumentChunk(
                    document_id=document.id,
                    chunk_index=i,
                    content=chunk.content,
                    section_title=chunk.section_title,
                    section_number=chunk.section_number,
                    page_number=chunk.page_number,
                    char_start=chunk.char_start,
                    char_end=chunk.char_end,
                    token_count=chunk.token_count,
                    embedding=embedding,
                    chunk_metadata=chunk.metadata
                )
                self.db.add(doc_chunk)
            
            document.chunk_count = len(all_chunks)
            document.ingestion_status = "completed"
            
        except Exception as e:
            document.ingestion_status = "failed"
            document.ingestion_error = str(e)
            raise
        
        finally:
            self.db.commit()
        
        return document
    
    def _detect_mime_type(self, file_path: str) -> str:
        """Detect MIME type from file extension."""
        ext = Path(file_path).suffix.lower()
        mime_map = {
            '.pdf': 'application/pdf',
            '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            '.txt': 'text/plain',
            '.md': 'text/markdown'
        }
        return mime_map.get(ext, 'application/octet-stream')
    
    def _calculate_file_hash(self, file_path: str) -> str:
        """Calculate SHA256 hash of file."""
        sha256 = hashlib.sha256()
        with open(file_path, 'rb') as f:
            for chunk in iter(lambda: f.read(8192), b''):
                sha256.update(chunk)
        return sha256.hexdigest()


def ingest_sample_documents(db: Session) -> List[Document]:
    """Ingest sample documents for testing."""
    pipeline = DocumentIngestionPipeline(db)
    
    # Create sample text documents since we don't have actual PDFs
    sample_docs = [
        {
            "title": "GDPR Article 5 - Principles of Processing",
            "source": "REG_GDPR",
            "document_type": "regulation",
            "content": """Article 5 - Principles relating to processing of personal data

1. Personal data shall be:
(a) processed lawfully, fairly and in a transparent manner in relation to the data subject ('lawfulness, fairness and transparency');
(b) collected for specified, explicit and legitimate purposes and not further processed in a manner that is incompatible with those purposes ('purpose limitation');
(c) adequate, relevant and limited to what is necessary in relation to the purposes for which they are processed ('data minimisation');
(d) accurate and, where necessary, kept up to date; every reasonable step must be taken to ensure that personal data that are inaccurate, having regard to the purposes for which they are processed, are erased or rectified without delay ('accuracy');
(e) kept in a form which permits identification of data subjects for no longer than is necessary for the purposes for which the personal data are processed ('storage limitation');
(f) processed in a manner that ensures appropriate security of the personal data, including protection against unauthorised or unlawful processing and against accidental loss, destruction or damage, using appropriate technical or organisational measures ('integrity and confidentiality').

2. The controller shall be responsible for, and be able to demonstrate compliance with, paragraph 1 ('accountability').""",
            "publication_date": datetime(2016, 4, 27)
        },
        {
            "title": "SOX Section 404 - Management Assessment of Internal Controls",
            "source": "REG_SOX",
            "document_type": "regulation",
            "content": """Section 404 - Management Assessment of Internal Controls

(a) RULES REQUIRED- The Commission shall prescribe rules requiring each annual report required by section 13(a) or 15(d) of the Securities Exchange Act of 1934 to contain an internal control report, which shall--
(1) state the responsibility of management for establishing and maintaining an adequate internal control structure and procedures for financial reporting; and
(2) contain an assessment, as of the end of the most recent fiscal year of the issuer, of the effectiveness of the internal control structure and procedures of the issuer for financial reporting.

(b) INTERNAL CONTROL EVALUATION AND REPORTING- With respect to the internal control assessment required by subsection (a), each registered public accounting firm that prepares or issues the audit report for the issuer shall attest to, and report on, the assessment made by the management of the issuer. An attestation made under this subsection shall be made in accordance with standards for attestation engagements issued or adopted by the Board. Any such attestation shall not be the subject of a separate engagement.""",
            "publication_date": datetime(2002, 7, 30)
        },
        {
            "title": "Data Protection Policy v2.3",
            "source": "POL_001",
            "document_type": "policy",
            "content": """Data Protection Policy

Version: 2.3
Effective Date: 2024-01-15
Owner: Data Protection Officer
Department: Legal & Compliance

1. Purpose
This policy establishes the framework for protecting personal data processed by the organization in accordance with GDPR, CCPA, and other applicable data protection regulations.

2. Scope
This policy applies to all employees, contractors, and third parties who process personal data on behalf of the organization.

3. Principles
3.1 Lawfulness, Fairness, and Transparency
All processing of personal data must have a lawful basis and be transparent to data subjects.

3.2 Purpose Limitation
Personal data must be collected for specified, explicit, and legitimate purposes.

3.3 Data Minimisation
Only personal data necessary for the stated purpose shall be collected and processed.

3.4 Accuracy
Personal data must be accurate and kept up to date.

3.5 Storage Limitation
Personal data shall not be kept longer than necessary.

3.6 Integrity and Confidentiality
Appropriate technical and organisational measures shall protect personal data.

4. Data Subject Rights
The organization shall facilitate the exercise of data subject rights including access, rectification, erasure, restriction, portability, and objection.

5. Data Breach Notification
Any personal data breach shall be reported to the DPO within 24 hours and to the supervisory authority within 72 hours where feasible.

6. Roles and Responsibilities
6.1 Data Protection Officer: Oversees compliance, conducts DPIAs, acts as contact point for supervisory authorities.
6.2 Data Controllers: Determine purposes and means of processing.
6.3 Data Processors: Process data on behalf of controllers.

7. Training
All staff processing personal data must complete annual data protection training.

8. Audit and Monitoring
Regular audits shall be conducted to verify compliance with this policy.""",
            "publication_date": datetime(2024, 1, 15)
        },
        {
            "title": "Access Control Policy v1.8",
            "source": "POL_002",
            "document_type": "policy",
            "content": """Access Control Policy

Version: 1.8
Effective Date: 2023-11-01
Owner: Chief Information Security Officer
Department: Information Technology

1. Purpose
Define requirements for controlling access to information systems and data.

2. Policy Statements
2.1 All access to systems shall be granted on a least-privilege basis.
2.2 Multi-factor authentication (MFA) is required for all remote access and privileged accounts.
2.3 Access rights shall be reviewed quarterly by system owners.
2.4 Access shall be revoked within 24 hours of termination or role change.
2.5 Generic/shared accounts are prohibited except for approved service accounts.
2.6 All access attempts shall be logged and monitored.

3. User Access Management
3.1 Provisioning: Access requests require manager approval and data owner sign-off.
3.2 Recertification: Quarterly access recertification campaigns for all systems.
3.3 Privileged Access: Privileged accounts require additional approval and monitoring.

4. System Access Controls
4.1 Password complexity: minimum 12 characters, complexity requirements.
4.2 Session timeout: 15 minutes of inactivity.
4.3 Failed login lockout: 5 attempts triggers 30-minute lockout.

5. Network Access Control
5.1 Network segmentation with firewall rules.
5.2 VPN required for remote access with MFA.
5.3 Wireless access controlled via enterprise authentication.

6. Compliance
Non-compliance may result in disciplinary action. Violations shall be reported via incident management process.""",
            "publication_date": datetime(2023, 11, 1)
        }
    ]
    
    documents = []
    for doc_data in sample_docs:
        # Write to temp file
        content = doc_data.pop("content")
        temp_path = f"/tmp/controllens_{doc_data['title'].lower().replace(' ', '_')}.txt"
        with open(temp_path, 'w') as f:
            f.write(content)
        
        doc_data["mime_type"] = "text/plain"
        doc = pipeline.ingest(temp_path, **doc_data)
        documents.append(doc)
        
        # Cleanup
        os.remove(temp_path)
    
    return documents